# -*- coding: utf-8 -*-
"""
群星（Stellaris）"进阶权威政体"数据提取脚本
================================================
从游戏文件中提取所有 advanced_authority_swap（进阶政体）的：
  · 名称（中文/英文）
  · 描述
  · 效果加成（modifier 解析为中文）
  · 特殊机制（tooltip 文本）
  · 触发条件（所属飞升路线）

输出：同级目录下生成 进阶政体数据.json

可选功能（需要额外依赖）：
  python 提取进阶政体.py --icons   # 转换政体图标为 PNG（需要 ImageMagick）
  python 提取进阶政体.py --excel   # 生成 Excel 汇总表（需要 openpyxl）

依赖：Python 3.8+（标准库即可完成核心提取）
说明：游戏目录如不同，修改下面的 GAME_DIR 变量。
"""

import re
import os
import sys
import json
import argparse

# ===================== 配置 =====================
GAME_DIR = r"D:\SteamLibrary\steamapps\common\Stellaris"
LANG_CN = "simp_chinese"          # 中文本地化目录名
LANG_EN = "english"               # 英文本地化目录名
OUT_JSON = "进阶政体数据.json"
OUT_ICON_DIR = "icons"
OUT_XLSX = "进阶政体汇总.xlsx"

AUTH_FILE = os.path.join(GAME_DIR, "common", "governments", "authorities", "00_authorities.txt")
LOC_DIR = os.path.join(GAME_DIR, "localisation")
ICON_SRC = os.path.join(GAME_DIR, "gfx", "interface", "icons", "governments", "authorities")


# ===================== 基础工具 =====================

def find_block(text, start):
    """从 start（'{' 之后的位置）匹配到配对的 '}'，返回块内容。"""
    depth = 1
    i = start
    while i < len(text) and depth > 0:
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
        i += 1
    return text[start:i - 1]


def load_loc(lang):
    """加载某个语言的全部本地化键值。"""
    loc = {}
    folder = os.path.join(LOC_DIR, lang)
    if not os.path.isdir(folder):
        return loc
    for fn in os.listdir(folder):
        if not fn.endswith(".yml"):
            continue
        try:
            with open(os.path.join(folder, fn), "r", encoding="utf-8-sig") as f:
                text = f.read()
        except Exception:
            continue
        for line in text.splitlines():
            m = re.match(r'\s+([\w_]+):\s*"(.*)"\s*(?:#.*)?$', line)
            if m:
                loc[m.group(1)] = m.group(2)
    return loc


def resolve_refs(s, loc):
    """展开 $key$ / $key$ 形式的本地化引用（最多递归 10 层）。"""
    n = 0
    while "$" in s and n < 10:
        m = re.search(r"\$([\w_]+)\$", s)
        if not m:
            break
        s = s.replace(m.group(0), loc.get(m.group(1), m.group(1)))
        n += 1
    return s


def clean_text(s):
    """清理群星富文本标记：颜色码、图标、变量、方括号引用等。"""
    if not s:
        return s
    s = (s.replace("\\n", " ")
          .replace("$t$", " ")
          .replace("$NEW_LINE$", " ")
          .replace("$TABBED_NEW_LINE$", " "))
    s = re.sub(r"§[A-Za-z]|§!", "", s)          # 颜色码 §Y §G §!
    s = re.sub(r"£[\w_]+£", "", s)              # 图标 £pop£
    s = re.sub(r"\$@?[\w_|%]+\$", "", s)        # 变量 $VAR$ / $@x|1%$
    s = re.sub(r"\[[^\]]*\]", "", s)            # 方括号引用 [concept_x]
    s = re.sub(r"\s{2,}", " ", s)
    return s.strip(" ：;；,，.-—")


# ===================== modifier 中文映射 =====================
# 未在表内的 modifier 键会保留英文原名（可自行补充）
MOD_CN = {
    # ---- 通用 / 产出类 ----
    "all_technology_research_speed": "全科研速度",
    "country_trade_produces_mult": "国家贸易产出",
    "country_unity_produces_mult": "国家凝聚力产出",
    "country_produces_mult": "国家产出",
    "pop_bonus_workforce_mult": "人口额外劳动力",
    "pop_political_power": "人口政治权力",
    "pop_amenities_usage_mult": "人口舒适度使用",
    "pop_environment_tolerance": "环境容忍度",
    "pop_ethics_shift_speed_mult": "思潮转变速度",
    "pop_government_ethic_attraction": "政府思潮吸引力",
    "pop_lifestyle_trade_value_add": "生活方式贸易价值",
    "pop_category_rulers_upkeep_mult": "统治者维护费",
    "planet_pop_assembly_add": "人口组装",
    "planet_resettlement_unemployed_mult": "失业人口迁移",
    "planet_colony_development_speed_mult": "殖民地发展速度",
    "colony_start_urban_districts_add": "初始城市区划",
    "auto_mod_monthly_mult": "自动改造进度",
    "bonus_pop_growth_mult": "人口增长加成",
    "modify_species_cost_mult": "物种改造费用",
    "subspecies_integration_mult": "亚种整合",
    "slave_market_cost_mult": "奴隶市场费用",
    "species_leader_exp_gain": "物种领袖经验",
    # ---- 建筑/设施 ----
    "planet_buildings_cost_mult": "建筑建造费用",
    "planet_structures_cost_mult": "设施建造费用",
    "planet_structures_upkeep_mult": "设施维护费",
    "planet_bureaucrats_produces_mult": "官僚产出",
    "planet_bureaucrats_unity_produces_mult": "官僚凝聚力产出",
    "planet_bureaucrats_upkeep_mult": "官僚维护费",
    "planet_politicians_unity_produces_add": "政客凝聚力产出",
    "planet_enforcers_produces_mult": "执法者产出",
    "planet_enforcers_unity_produces_add": "执法者凝聚力产出",
    "planet_enforcers_engineering_research_produces_add": "执法者工程研究产出",
    "job_enforcer_stability_add": "执法者稳定度",
    "job_soldier_stability_add": "士兵稳定度",
    "station_researchers_produces_mult": "研究站产出",
    # ---- 人口阶级 ----
    "pop_cat_worker_bonus_workforce_mult": "劳工额外劳动力",
    "pop_cat_specialist_bonus_workforce_mult": "专家额外劳动力",
    "pop_cat_ruler_bonus_workforce_mult": "统治者额外劳动力",
    "pop_cat_simple_drone_bonus_workforce_mult": "简易无人机额外劳动力",
    "pop_cat_complex_drone_bonus_workforce_mult": "复合无人机额外劳动力",
    "pop_cat_purge_bonus_workforce_mult": "净化人口额外劳动力",
    "pop_cat_worker_political_power": "劳工政治权力",
    "pop_cat_specialist_political_power": "专家政治权力",
    "pop_cat_ruler_political_power": "统治者政治权力",
    "cyborg_pop_bonus_workforce_mult": "义体人口额外劳动力",
    "planet_pops_cyborgs_upkeep_mult": "义体人口维护费",
    # ---- 议会 / 领袖 ----
    "council_agenda_cooldown_mult": "议事日程冷却",
    "council_agenda_early_activation_cost_mult": "议事日程提前激活费用",
    "council_agenda_finishing_effect_duration_mult": "议事日程完成效果时长",
    "council_agenda_progress_speed": "议事日程进度",
    "councilor_exp_gain": "议员经验获取",
    "councilor_ruler_dictatorial_exp_gain": "独裁统治者经验获取",
    "councilor_skill_add": "议员技能",
    "leader_initial_skill": "领袖初始技能",
    "leader_skill_add": "领袖技能",
    "negative_traits_leader": "领袖负面特质",
    "election_term_years_add": "选举任期",
    # ---- 舰船 / 军事 ----
    "ship_fire_rate_mult": "舰船开火率",
    "ship_speed_mult": "舰船速度",
    "ship_armor_regen_add_perc": "舰船装甲回复",
    "ship_hull_regen_add_perc": "舰船船体回复",
    "starbase_shipyard_build_speed_mult": "船坞建造速度",
    # ---- 外交 / 帝国 ----
    "diplo_weight_economy_mult": "外交权重（经济）",
    "country_branch_office_influence_cost_mult": "分部影响力费用",
    "operations_influence_cost_mult": "行动影响力费用",
    "empire_size_psionic_pops_mult": "帝国规模（灵能人口）",
    "edicts_cost_mult": "法令费用",
    # ---- 机械 ----
    "MACHINE_species_trait_picks_add": "机械物种特质点数",
    "faction_approval": "派系支持度",
    "country_leader_pool_size": "领袖池",
    "pop_ethic_spiritualist_attraction_mult": "唯心思潮吸引力",
    "pop_cyborg_happiness": "义体人幸福度",
    "pop_non_cyborg_happiness": "非义体人幸福度",
    "pop_factions_produces_mult": "派系产出",
    "leaders_upkeep_mult": "领袖维护费",
    "leaders_cost_mult": "领袖花费",
    "empire_size_pops_mult": "帝国规模（人口）",
    "empire_size_penalty_mult": "帝国规模惩罚",
    "empire_size_districts_mult": "帝国规模（区划）",
    "empire_size_colonies_mult": "帝国规模（殖民地）",
    "empire_size_systems_mult": "帝国规模（星系）",
    "empire_size_mult": "帝国规模",
    "empire_size_branch_office_mult": "帝国规模（分部）",
    "planet_jobs_specialist_produces_mult": "专家产出",
    "planet_jobs_worker_produces_mult": "劳工产出",
    "planet_jobs_simple_drone_produces_mult": "简易无人机产出",
    "planet_jobs_complex_drone_produces_mult": "复合无人机产出",
    "planet_jobs_trade_produces_mult": "贸易岗位产出",
    "planet_jobs_unity_produces_mult": "凝聚力岗位产出",
    "planet_traders_trade_produces_add": "商人贸易产出",
    "planet_bureaucrats_unity_produces_add": "官僚凝聚力产出",
    "planet_bureaucrats_trade_produces_add": "官僚贸易产出",
    "planet_bureaucrats_physics_research_produces_add": "官僚物理研究产出",
    "planet_crime_add": "犯罪率",
    "planet_crime_no_happiness_mult": "犯罪率（无幸福度）",
    "planet_stability_add": "稳定度",
    "planet_stability_mult": "稳定度%",
    "pop_happiness": "人口幸福度",
    "all_technology_research_speed": "全科研速度",
    "country_edict_fund_add": "法令资金",
    "country_edict_fund_mult": "法令资金%",
    "edicts_upkeep_mult": "法令维护费",
    "country_naval_cap_mult": "海军容量",
    "country_naval_capacity_contribution_from_subjects_mult": "附庸海军贡献",
    "country_enclave_capacity_add": "飞地容量",
    "country_starbase_capacity_add": "恒星基地容量",
    "starbase_buildings_trade_produces_add": "恒星基地建筑贸易产出",
    "starbase_shipyard_build_cost_mult": "船坞建造费用",
    "ship_cloaking_strength_add": "舰船隐身强度",
    "army_damage_mult": "陆军伤害",
    "country_war_exhaustion_mult": "厌战度",
    "election_term_years_mult": "选举任期",
    "pop_ethic_egalitarian_attraction_mult": "平等思潮吸引力",
    "intel_encryption_add": "情报加密",
    "intel_decryption_add": "情报解密",
    "refugee_attraction": "难民吸引力",
    "founder_species_growth_mult": "主体物种增长",
    "pop_citizen_happiness": "公民幸福度",
    "envoy_improve_relations_mult": "特使改善关系",
    "envoy_harm_relations_mult": "特使恶化关系",
    "envoy_attunement_mult": "特使协调",
    "envoys_add": "特使数量",
    "leaders_skill_add": "领袖技能",
    "ruler_skill_add": "统治者技能",
    "country_leader_cap_add": "领袖上限",
    "num_tech_alternatives_add": "科研备选项",
    "country_scientist_cap_add": "科学家上限",
    "country_official_cap_add": "官员上限",
    "psionic_aura_hyperlane_range_add": "灵能光环超空间航程",
    "psionic_aura_effectiveness_mult": "灵能光环效果",
    "psionic_pop_bonus_workforce_mult": "灵能人口额外劳动力",
    "leaders_lifespan_add": "领袖寿命",
    "leader_lifespan_add": "领袖寿命",
    "leaders_initial_skill_add": "领袖初始技能",
    "leader_recruitment_age": "领袖招募年龄",
    "logistic_growth_mult": "后勤增长",
    "planet_pop_assembly_mult": "人口组装",
    "planet_researchers_upkeep_mult": "研究员维护费",
    "planet_buildings_upkeep_mult": "建筑维护费",
    "country_subject_acceptance_add": "附庸接受度",
    "divided_patrongage_max_subjects": "最大附庸数",
    "starbase_buildings_trade_produces_add": "恒星基地建筑贸易产出",
    "branch_office_value_mult": "分部价值",
    "country_life_style_trade_mult": "生活方式贸易",
    "traits_points_add": "物种特质点数",
    "pop_cat_worker_produces_mult": "劳工阶级产出",
    "pop_cat_specialist_produces_mult": "专家阶级产出",
    "pop_cat_ruler_produces_mult": "统治者阶级产出",
    "pop_cat_worker_happiness": "劳工阶级幸福度",
    "pop_cat_specialist_happiness": "专家阶级幸福度",
    "pop_cat_ruler_happiness": "统治者阶级幸福度",
    "country_produces_mult": "国家产出",
}

MOD_VALUE_SUFFIX = {  # 后缀说明（可留空）
}


SKIP_KEYS = {"show_only_custom_tooltip", "hidden_effect", "custom_tooltip"}


def load_scripted_vars(text):
    """收集 authorities 文件里的 @变量 定义（如 @cyber_collectivist_stability = 5）。"""
    vars_ = {}
    for m in re.finditer(r"^@(\w+)\s*=\s*(-?[\d.]+)\s*$", text, re.M):
        vars_[m.group(1)] = float(m.group(2))
    return vars_


def fmt_modifier(raw, vars_=None):
    """把 modifier 块文本解析为 '中文名: +值' 列表。支持 @变量 引用。"""
    vars_ = vars_ or {}
    out = []
    if not raw:
        return out
    for line in raw.splitlines():
        line = line.strip()
        if not line or line.startswith("#") or line.startswith("{"):
            continue
        # 跳过配置指令（show_only_custom_tooltip 等）
        key0 = line.split("=", 1)[0].strip()
        if key0 in SKIP_KEYS:
            continue
        m = re.match(r"^([\w_]+)\s*=\s*(@[\w_]+|-?[\d.]+)\s*$", line)
        if not m:
            out.append(line)
            continue
        key, raw_val = m.group(1), m.group(2)
        if raw_val.startswith("@"):
            val = vars_.get(raw_val[1:])
            if val is None:
                continue  # 变量未定义则跳过该行
        else:
            val = float(raw_val)
        name = MOD_CN.get(key, key)
        sign = "+" if val >= 0 else ""
        # 百分比类（|v|<1 且键含 mult/attraction/speed 等特征）显示为 %
        if -1 < val < 1 and val != 0 and ("mult" in key or "attraction" in key or "speed" in key):
            shown = f"{sign}{val*100:.1f}%"
        else:
            shown = f"{sign}{val:g}"
        out.append(f"{name}: {shown}")
    return out


# ===================== 主提取逻辑 =====================

def extract_swaps():
    """解析 00_authorities.txt 中所有 advanced_authority_swap 块。"""
    with open(AUTH_FILE, "r", encoding="utf-8-sig") as f:
        text = f.read()

    # 1) 找到每个顶层 authority 块
    auths = {}
    for m in re.finditer(r"^(auth_\w+)\s*=\s*\{", text, re.M):
        auths[m.group(1)] = find_block(text, m.end())

    # 2) 在每个 authority 块中找 advanced_authority_swap
    swaps = []
    for auth, block in auths.items():
        for m in re.finditer(r"advanced_authority_swap\s*=\s*\{", block):
            sw = find_block(block, m.end())
            name_m = re.search(r'name\s*=\s*"([^"]+)"', sw)
            desc_m = re.search(r'description\s*=\s*"([^"]+)"', sw)
            mod_m = re.search(r"modifier\s*=\s*\{(.*?)\n\s*\}", sw, re.S)
            trig_m = re.search(r"trigger\s*=\s*\{(.*?)\n\s*\}", sw, re.S)
            if name_m:
                swaps.append({
                    "auth": auth,                       # 基座权威（auth_democratic 等）
                    "key": name_m.group(1),             # 政体键（auth_cyber_... 等）
                    "desc_key": desc_m.group(1) if desc_m else "",
                    "modifier": mod_m.group(1).strip() if mod_m else "",
                    "trigger": trig_m.group(1).strip() if trig_m else "",
                })
    return swaps


def build_data():
    """组装完整数据：名称 / 描述 / 效果 / 特殊机制 / 触发条件。"""
    cn = load_loc(LANG_CN)
    en = load_loc(LANG_EN)
    with open(AUTH_FILE, "r", encoding="utf-8-sig") as f:
        vars_ = load_scripted_vars(f.read())
    swaps = extract_swaps()
    data = []
    for s in swaps:
        key = resolve_refs(s["key"], cn)
        name_cn = cn.get(key, key)
        name_en = en.get(key, "")
        desc = clean_text(resolve_refs(cn.get(s["desc_key"], "") if s["desc_key"] else "", cn))
        mods = fmt_modifier(s["modifier"], vars_)
        tt_key = key + "_tt"
        tt = clean_text(resolve_refs(cn.get(tt_key, ""), cn))
        data.append({
            "key": key,
            "base_authority": s["auth"],
            "name": name_cn,
            "name_en": name_en,
            "desc": desc,
            "mods": mods,
            "tt": tt,
            "trigger": s["trigger"],
        })
    return data


# ===================== 可选功能 =====================

def export_icons(data):
    """把政体图标从 dds 转换为 png（需要已安装 ImageMagick，命令为 magick）。"""
    import subprocess
    os.makedirs(OUT_ICON_DIR, exist_ok=True)
    ok, fail = 0, []
    for d in data:
        src = os.path.join(ICON_SRC, d["key"] + ".dds")
        dst = os.path.join(OUT_ICON_DIR, d["key"] + ".png")
        if not os.path.exists(src):
            fail.append(d["key"])
            continue
        r = subprocess.run(["magick", src, "-resize", "96x96", dst], capture_output=True)
        if r.returncode == 0:
            ok += 1
        else:
            fail.append(d["key"])
    print(f"图标导出: 成功 {ok}, 失败 {len(fail)}")


def export_excel(data):
    """导出 Excel 汇总表（需要 openpyxl）。"""
    import openpyxl
    from openpyxl.styles import Font, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "进阶政体"
    headers = ["序号", "政体名称", "英文名", "基座权威", "描述", "效果加成", "特殊机制", "我的评分"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, d in enumerate(data, 1):
        ws.append([
            i, d["name"], d["name_en"], d["base_authority"],
            d["desc"], "\n".join(d["mods"]), d["tt"], None,
        ])
    for col, w in zip("ABCDEFGH", [6, 24, 28, 12, 40, 44, 50, 10]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(OUT_XLSX)
    print(f"Excel 已导出: {OUT_XLSX}")


# ===================== 入口 =====================

def main():
    parser = argparse.ArgumentParser(description="提取群星进阶权威政体数据")
    parser.add_argument("--icons", action="store_true", help="导出政体图标 PNG")
    parser.add_argument("--excel", action="store_true", help="导出 Excel 汇总表")
    args = parser.parse_args()

    print(f"读取游戏目录: {GAME_DIR}")
    if not os.path.exists(AUTH_FILE):
        print("错误：找不到 authorities 文件，请修改脚本顶部的 GAME_DIR。")
        sys.exit(1)

    data = build_data()
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"已提取 {len(data)} 个进阶政体 -> {OUT_JSON}")

    if args.icons:
        export_icons(data)
    if args.excel:
        export_excel(data)

    # 简要预览
    for d in data[:5]:
        print(f"  {d['name']:<12} | 基座: {d['base_authority']} | 效果 {len(d['mods'])} 条")


if __name__ == "__main__":
    main()
